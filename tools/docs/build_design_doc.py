from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
ASSETS = ROOT / "docs" / "assets"
OUTPUTS = ROOT / "docs"
OUTPUTS.mkdir(parents=True, exist_ok=True)
OUT = OUTPUTS / "CampusNetGuard开发设计与实现说明.docx"

ACCENT = "245CB8"
ACCENT_LIGHT = "EAF2FF"
ACCENT_MID = "DCEBFF"
GRAY = RGBColor(80, 90, 105)
DARK = RGBColor(25, 44, 72)

def set_run_font(run, name="Microsoft YaHei", size=None, bold=None, color=None):
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), name)

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + tag))
        if node is None:
            node = OxmlElement("w:" + tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_cell_text(cell, text, bold=False, size=9.5, color=None, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(str(text))
    set_run_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)

def style_table(table, header_fill=ACCENT_LIGHT):
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            if row_idx == 0:
                set_cell_shading(cell, header_fill)
            else:
                set_cell_shading(cell, "FFFFFF")

def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)

def set_col_widths(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)

def add_table(doc, headers, rows, widths=None, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    style_table(table)
    for idx, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], h, bold=True, size=font_size)
    repeat_header(table.rows[0])
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            set_cell_text(cells[idx], value, size=font_size)
    if widths:
        set_col_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run_font(r, size=9, color=GRAY)

def add_image(doc, filename, caption, width=6.35):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(ASSETS / filename), width=Inches(width))
    add_caption(doc, caption)

def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, ACCENT_LIGHT)
    set_cell_margins(cell, top=140, start=170, bottom=140, end=170)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_run_font(r, size=11, bold=True, color=DARK)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run_font(r, size=10.5)

def add_para(doc, text, keep_with_next=False):
    p = doc.add_paragraph()
    if keep_with_next:
        p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size=10.5)
    return p

def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 12.5, bold=True, color=DARK)
    return p

def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fld_end])
    set_run_font(run, size=9, color=GRAY)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal.font.size = Pt(10.5)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.paragraph_format.line_spacing = 1.22
normal.paragraph_format.space_after = Pt(5)

for style_name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 12.5), ("List Bullet", 10.5)):
    st = styles[style_name]
    st.font.name = "Microsoft YaHei"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(size)
    st.font.color.rgb = DARK if style_name != "Title" else RGBColor(0, 0, 0)

header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = hp.add_run("校园网断线自动认证助手  |  开发设计与实现总结")
set_run_font(hr, size=8.5, color=GRAY)

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run("第 ")
set_run_font(fr, size=9, color=GRAY)
add_page_number(fp)
fr2 = fp.add_run(" 页")
set_run_font(fr2, size=9, color=GRAY)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(80)
title.paragraph_format.space_after = Pt(12)
tr = title.add_run("校园网断线自动认证助手设计与开发总结")
set_run_font(tr, size=24, bold=True, color=RGBColor(0, 0, 0))

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.paragraph_format.space_after = Pt(26)
sr = sub.add_run("Windows 10/11、.NET Framework 与 WinForms 实现思路")
set_run_font(sr, size=12, color=GRAY)

meta = doc.add_table(rows=4, cols=2)
style_table(meta)
meta_rows = [
    ("文档用途", "供研究程序的设计思路、实现分层和开发迭代过程"),
    ("目标平台", "Windows 10/11，.NET Framework 4.8，WinForms 托盘程序"),
    ("核心目标", "校园网断线后自动检测并重新认证；有线物理断开时尝试 Wi-Fi 回退"),
    ("文档版本", "2026-09-29，精简版工程总结"),
]
for i, (k, v) in enumerate(meta_rows):
    set_cell_text(meta.rows[i].cells[0], k, bold=True)
    set_cell_text(meta.rows[i].cells[1], v)
set_col_widths(meta, [1.35, 5.0])
doc.add_page_break()

add_heading(doc, "1. 文档定位与结论", 1)
add_para(doc, "这份文档不是逐行代码说明，也不替代使用手册，而是把程序为什么这样设计、由哪些部分组成、关键判断如何工作、开发中遇到哪些问题以及最后如何验证，压缩成一条可追踪的工程思路。")
add_callout(doc, "一句话结论", "程序采用“轻量后台监控 + 状态机 + 门户认证 + 浏览器回退”的方案：网络正常时几乎不做事；连续检测到断网后才访问校园网门户；普通认证失败每 30 秒持续重试，浏览器回退受独立节流；学校计划断网时段和手动暂停会停止自动操作，但不会主动断开网络。")
add_para(doc, "最终实现为单个可执行文件，默认不需要管理员权限，不使用 Windows 服务。程序只在当前用户登录后运行，通过托盘界面提供自动重连开关、立即检测、账号设置、日志位置和状态查看等功能。")

add_heading(doc, "2. 问题背景与可行性判断", 1)
add_para(doc, "目标电脑需要通过 UU 远程等工具长期保持可访问。校园网认证掉线后，远程端无法打开门户页面，也就无法替电脑完成手动登录。因此，问题的本质不是“自动打开网页”，而是让一个本地程序在用户不在电脑旁时完成以下闭环：识别断线、确认是否需要认证、提交认证、验证是否真正恢复，并在失败时避免盲目重试。")
add_para(doc, "从技术上判断，这一目标可行。Windows 提供网络变化事件和网卡信息，程序可以识别当前物理有线/无线连接；校园网门户使用普通 HTTP 表单，程序可以动态解析表单字段并提交。主要风险在于门户页面结构可能变化、直接访问门户可能返回空页面，以及虚拟网卡或 VPN 会让“看起来在线”与“实际不能上网”不一致。")
add_table(doc, ["设计问题", "采用的判断"], [
    ("能否检测断网", "能。使用联网探测内容校验，而不是只看网卡状态"),
    ("能否自动认证", "能。动态解析门户表单并提交账号密码；必要时回退到本机 Chrome/Edge"),
    ("会不会误判", "通过连续两次探测、绑定物理网卡 IP、过滤虚拟网卡降低概率"),
    ("会不会无限重试", "普通认证失败每 30 秒持续重试，浏览器回退最多每 2 分钟一次"),
    ("密码如何保存", "当前 Windows 用户 DPAPI 加密，不写死在程序中"),
], widths=[1.75, 4.6])
add_para(doc, "需要明确的边界是：程序不能唤醒已关机、睡眠或注销的电脑，也不能保证在无线网卡被硬件开关关闭时自行恢复；门户认证规则变化时，仍需要人工确认并更新程序。")

add_heading(doc, "3. 设计原则与取舍", 1)
add_bullets(doc, [
    "低干扰：仅在网络事件发生或 30 秒兜底检查时工作，不在后台持续刷网页。",
    "先验证再认证：先判断互联网是否真的不可用，再决定是否提交门户登录。",
    "安全优先：不硬编码账号密码，默认不请求管理员权限，不记录密码、Cookie 或完整认证数据。",
    "失败可控：认证异常时记录原因并进入等待，不把一次页面解析失败当成无限重试的理由。",
    "可手动接管：托盘开关、立即检测、暂停和有线恢复策略都可见、可操作。",
])
add_para(doc, "这些取舍决定了程序不会追求“任何情况下都自动恢复”。当出现明确账号错误、动态验证码、无线网卡关闭或门户结构未知时，程序更倾向于安全停止并提示用户，而不是反复提交。")

add_heading(doc, "4. 总体架构与模块划分", 1)
add_para(doc, "程序以 GuardContext 作为调度中心，把网络状态检测、门户认证、浏览器回退、配置存储和托盘界面连接起来。网络变化事件负责快速触发，30 秒定时器负责兜底，二者都进入同一个状态机，避免多套逻辑互相冲突。")
add_image(doc, "architecture.png", "图 1  程序总体架构：事件驱动、状态机、认证与回退模块", width=6.25)
add_table(doc, ["模块", "主要职责"], [
    ("Program / GuardContext", "启动、单实例、状态机、30 秒重试和暂停调度"),
    ("NetworkUtils", "识别真实有线/无线网卡、探测互联网、发现校园网环境"),
    ("PortalAuthenticator", "读取门户页面、解析表单、提交认证、解析响应"),
    ("BrowserPortalFallback", "直接请求失败时用本机 Chrome/Edge 执行门户脚本"),
    ("WifiManager", "检测无线网卡、连接已保存的 gdei-WiFi-S 配置"),
    ("Models / AppPaths", "保存设置、DPAPI 密码、日志目录和运行参数"),
    ("Ui", "托盘菜单、状态窗口、账号设置和日志位置选择"),
], widths=[1.75, 4.6])
add_para(doc, "代码文件和文档保持同一目录，便于排查：Program.cs 负责入口，GuardContext.cs 是主要流程，NetworkUtils.cs 和 PortalAuthenticator.cs 分别承担网络与门户逻辑，BrowserPortalFallback.cs 是兜底路径，Ui.cs 负责界面。")

add_heading(doc, "5. 断线检测与运行流程", 1)
add_para(doc, "程序只把物理以太网和 Wireless80211 无线网卡当作候选连接。包含 tap、vpn、virtual、zerotier、radmin、vmware、virtualbox、wintun、wireguard 等关键词的适配器会被忽略，避免 UU、ZeroTier 或其他虚拟网卡在线时掩盖真实校园网断线。")
add_para(doc, "互联网探测使用 http://www.msftconnecttest.com/connecttest.txt，并校验返回内容是否包含 Microsoft Connect Test。第一次失败不会立即认证，而是约 5 秒后再测一次；连续两次失败才进入需要处理的状态。探测请求绑定当前物理网卡 IPv4，减少 VPN 默认路由造成的误判。")
add_bullets(doc, [
    "Windows NetworkChange 事件：网络地址或可用性变化时立即排队检查。",
    "30 秒兜底定时器：防止事件丢失，同时保持较低资源占用。",
    "两次失败判定：只在稳定断网后进入认证，短暂抖动不触发。",
    "认证后再次验证：提交成功后不能只相信页面文本，还要重新探测互联网。",
])
add_image(doc, "runtime_flow.png", "图 2  运行状态流程：在线、认证、Wi-Fi 回退、暂停和持续重试", width=6.25)
add_para(doc, "状态机大致经历：启动并加载配置，进入在线或等待网络；断网后判断这是否属于学校计划断网时段；若不是，则尝试有线门户认证；普通失败在 30 秒后继续尝试；物理断开时进入 Wi-Fi 回退。计划暂停、手动暂停、凭据错误和人工处理会阻止自动提交，但不会让程序退出托盘。")

add_heading(doc, "6. 校园网门户认证设计", 1)
add_para(doc, "有线门户使用给定地址 https://nac.gdei.edu.cn/portal/2021/login.jsp?nasip=192.168.104.2 作为兼容入口。Wi-Fi 门户地址不直接猜测 nasip，而是在连接 gdei-WiFi-S 并取得 IP 后，通过 HTTP 强制门户跳转捕获实际门户地址。")
add_para(doc, "直接认证流程尽量模拟用户浏览器：使用普通浏览器 User-Agent、保持 Cookie、启用 TLS 1.2、禁用代理；读取页面后动态解析 form1 中的 input 和 select，并保留 wlanacname、nasip 等隐藏字段。用户名写入 userName，密码按门户要求做 UTF-8 字节的 Base64 后写入 userPwd，再提交到 /portal/pws?t=li&ifEmailAuth=false。")
add_table(doc, ["步骤", "处理逻辑"], [
    ("1  获取页面", "请求门户地址并识别登录表单；页面为空或表单缺失时不盲目提交"),
    ("2  解析表单", "保留隐藏字段，识别 userName、userPwd 和动态参数"),
    ("3  提交认证", "按门户格式提交；提交内容在日志中不完整记录"),
    ("4  解析响应", "先 URL 解码，必要时再 Base64 解码，提取 errorNumber、portalLink、e_c、e_d"),
    ("5  判定结果", "errorNumber=1 视为成功；63032 或明确凭据错误进入停止重试状态"),
    ("6  回退方案", "直接请求失败时使用本机 Chrome/Edge 加载门户并尝试门户自带脚本或 fetch"),
], widths=[1.35, 5.0])
add_image(doc, "portal_flow.png", "图 3  门户认证主流程与浏览器回退流程", width=6.25)

add_heading(doc, "7. 有线物理断开与 Wi-Fi 回退", 1)
add_para(doc, "有线网卡拔出或物理连接变为 Unplugged 时，程序不会立刻切换 Wi-Fi，而是等待约 10 秒。短暂抖动后恢复有线时继续使用有线；持续断开才启动回退。这个延迟用于区分偶发接触不良和真实拔线。")
add_para(doc, "Wi-Fi 回退只连接 Windows 中已保存的 gdei-WiFi-S，不尝试 gdei-WiFi。有线持续断开后，程序先检查 WLAN 无线电状态；如果只是 Windows 软件 Wi-Fi 开关关闭，会通过 WLAN API 尝试开启。硬件开关关闭或网卡被设备管理器禁用时，程序不会请求管理员权限强行绕过，而是提示用户。为避免在家或其他网络环境误连，程序还要求最近有校园网上下文，或附近确实能看到该 SSID。")
add_table(doc, ["有线恢复策略", "行为"], [
    ("稳定联网 60 秒后切回", "默认策略。Wi-Fi 间歇可用时不立即切回，降低抖动"),
    ("保留 Wi-Fi，手动切回", "程序不再自动操作，用户可通过“立即恢复有线”执行"),
    ("保持 Wi-Fi，不自动恢复", "适合用户希望固定使用无线的情况"),
], widths=[1.7, 4.65])
add_para(doc, "如果是 Windows 软件 Wi-Fi 开关关闭，程序会尝试通过 WLAN API 自动开启；如果硬件开关关闭、网卡被设备管理器禁用或系统策略阻止，程序不会请求管理员权限强行绕过，而是通过托盘提示用户手动处理。这样保留了“普通权限即可运行”的设计前提。")

add_heading(doc, "8. 计划暂停、持续重试与防误判", 1)
add_para(doc, "校园网在周日到周四凌晨 01:00–06:00 计划断网。程序对这个时间段采用完全暂停：不探测、不认证、不切换 Wi-Fi、不发无意义通知。时间边界采用左闭右开，01:00:00 开始暂停，05:59:59 仍暂停，06:00:00 恢复并立即检查。周五和周六全天运行。")
add_table(doc, ["周期", "自动处理"], [
    ("周日、周一、周二、周三、周四 01:00–06:00", "暂停检测、认证、Wi-Fi 切换和提示"),
    ("上述日期的其他时间", "正常运行"),
    ("周五、周六", "全天正常运行，不进入计划暂停"),
], widths=[3.1, 3.25])
add_para(doc, "普通认证失败固定等待 30 秒后继续尝试，不再限制每小时 6 次。浏览器回退最多每 2 分钟启动一次，避免高频启动 Chrome 或 Edge。认证成功后会清空失败计数；账号或密码明确错误、页面出现可见且可提交的验证码或动态密码字段时，程序停止自动认证，等待用户修改凭据或手动处理。")
add_para(doc, "开发中特别避免了两类误判：第一，不把页面中隐藏的中文模板文字当成人机验证；第二，不把一次“门户页面为空或未返回登录表单”直接当成最终失败。当前逻辑会在异常后重新探测互联网，如果网络已恢复，则记录恢复并按成功处理。因此，日志中出现一条 WARN 不一定代表最终仍然断网。")

add_heading(doc, "9. 安全、配置、日志与界面", 1)
add_para(doc, "账号密码第一次由用户输入。密码只以当前 Windows 用户的 DPAPI 密文保存在设置中，程序不硬编码密码，也不在日志中写入密码、Cookie 或完整认证数据。日志中的账号会脱敏，例如只保留前两位并显示为 20****。")
add_table(doc, ["项目", "设计"], [
    ("设置", "%LOCALAPPDATA%\\CampusNetGuard\\settings.xml"),
    ("默认日志", "%LOCALAPPDATA%\\CampusNetGuard\\logs"),
    ("日志位置修改", "托盘菜单选择“日志保存位置”，可移动已有 app-*.log 文件"),
    ("开机启动", "写入 HKCU\\Software\\Microsoft\\Windows\\CurrentVersion\\Run"),
    ("单实例", "使用 Local\\CampusNetGuard.SingleInstance 防止重复启动"),
], widths=[1.55, 4.8])
add_para(doc, "托盘菜单包含：查看状态、自动重连、立即检测、有线恢复策略、账号设置、查看日志、日志保存位置、有线断开时尝试开启 Wi-Fi、开机启动和退出。状态窗口集中显示有线连接、Wi-Fi SSID、互联网状态、校园环境判断、计划暂停、自动开启 Wi-Fi 设置、最近认证结果、下次重试时间和恢复策略。手动暂停只停止自动操作，程序仍保留在托盘中。")
add_para(doc, "程序文件和配置、日志相互独立。移动 exe 或快捷方式不会自动改变已有设置和日志；如果使用快捷方式或开机启动，只需要确认快捷方式目标、启动项和实际 exe 路径一致。这样可以把程序放到其他目录，而配置仍保留在当前用户的应用数据目录中。")

add_heading(doc, "10. 开发过程与问题演进", 1)
add_para(doc, "开发不是一次性写完，而是围绕真实使用中的故障逐步收敛。最初的目标只是“断线后自动打开网页登录”，随后发现单一网络判断不足以覆盖校园网定时断网、虚拟网卡干扰和有线物理断开，于是逐步加入事件检测、状态机、暂停规则、Wi-Fi 回退和浏览器兜底。")
add_table(doc, ["观察到的现象", "设计上的修正"], [
    ("校园网断开后没有明显重试", "把网络变化事件与 30 秒兜底检查合并，并绑定物理网卡 IP"),
    ("门户返回空页面或没有表单", "保留直接认证作为主路径，增加 Chrome/Edge 浏览器回退"),
    ("日志先报失败，后来又能上网", "失败后再次探测互联网；恢复时按成功处理，WARN 只是一次尝试结果"),
    ("隐藏模板文字被误认为验证码", "只检查可见且可提交的验证码或动态密码字段"),
    ("害怕学校计划断网期间反复重试", "周日到周四 01:00–06:00 暂停自动操作，06:00 恢复；不会主动断开网络"),
    ("有线随机拔出后又恢复", "设置约 10 秒防抖，持续断开后才进入 Wi-Fi 回退"),
    ("需要控制日志体积或位置", "增加日志目录设置，并可迁移已有 app-*.log"),
], widths=[2.2, 4.15])
add_para(doc, "这些修正保持了同一个核心结构：状态机负责“现在该做什么”，网络模块负责“是否真的在线”，门户模块负责“如何完成认证”，浏览器回退负责“直接访问失败时再试一次”。这样新问题通常只需要增加判断或策略，不必推翻整体程序。")

add_heading(doc, "11. 测试、验收与资源占用", 1)
add_para(doc, "当前离线回归测试覆盖计划暂停时间边界、周五周六行为、30 秒重试计划、 Wi-Fi 无线电状态策略、网络恢复日志语义、校园网特征识别、门户表单解析、账号日志脱敏、日志目录修改、单实例、验证码误判、63032 凭据错误、响应解码和失败状态清理等场景。")
add_table(doc, ["验收方向", "结果"], [
    ("暂停边界", "周日到周四 01:00–06:00 正确暂停，06:00 恢复"),
    ("真实网卡过滤", "虚拟网卡存在时不会把它当成有效互联网连接"),
    ("门户解析", "隐藏字段保留，表单字段可动态识别"),
    ("安全存储与日志", "密码使用 DPAPI，账号脱敏，敏感信息不写入日志"),
    ("失败控制", "普通故障每 30 秒重试，凭据错误和人工验证要求停止自动提交"),
    ("真实现场", "真实拔线和真实认证只适合在用户在场时手动验证"),
], widths=[1.75, 4.6])
add_para(doc, "空闲时程序内存通常约 20–40 MB，CPU 接近 0；它依靠事件和 30 秒兜底定时器，不存在常驻浏览器或持续高负载进程。只有直接认证失败并进入浏览器回退时，Chrome/Edge 会临时增加约 100–300 MB 内存，用完即关闭并清理临时配置。")

add_heading(doc, "12. 部署、限制与维护要点", 1)
add_table(doc, ["要点", "说明"], [
    ("运行方式", "当前用户登录后运行；无需管理员权限，不创建服务"),
    ("运行位置", "可移动 exe，但移动后要同步快捷方式和开机启动目标"),
    ("配置与日志", "保存在 %LOCALAPPDATA%，通常不随 exe 移动"),
    ("无线前提", "软件 Wi-Fi 开关可尝试自动开启，硬件开关需可用；gdei-WiFi-S 已保存"),
    ("不能覆盖的场景", "关机、睡眠、注销、无线硬件关闭、门户规则根本变化"),
    ("维护重点", "门户地址、登录字段、回退地址和认证响应结构可能变化，需要保留日志并更新程序"),
], widths=[1.55, 4.8])
add_para(doc, "程序的安全策略是“高风险错误停住，普通故障持续恢复”。如果直接请求失败，会在节流允许时尝试浏览器回退；回退仍失败则每 30 秒继续。如果发现账号密码明显错误、出现人工验证要求，或认证响应无法解释，程序保留状态和日志，不进行高风险连续提交。")

add_heading(doc, "13. 结论与运行速查", 1)
add_para(doc, "这套程序的设计重点不是把门户页面硬编码成一个脚本，而是把“网络是否真的可用”和“是否值得尝试认证”分开判断，再用状态机协调有线、无线、暂停、重试节流和用户手动控制。这样既能处理校园网的日常掉线，也能在异常情况下保持可控和可排查。")
add_bullets(doc, [
    "正常上网：程序只做低频兜底检查和事件监听，不访问门户。",
    "确认断网：连续两次探测失败后再进入认证，减少误判。",
    "认证失败：普通错误每 30 秒重试；浏览器回退最多每 2 分钟一次。",
    "计划暂停：周日到周四 01:00–06:00 仅暂停自动操作，06:00 恢复，不会主动断网。",
    "Wi-Fi 回退：软件开关关闭时尝试自动开启，硬件开关关闭时提示用户。",
    "有线断开：约 10 秒后仍断开才尝试 gdei-WiFi-S。",
    "安全边界：账号密码本机加密保存，程序默认普通权限运行。",
])
final_p = doc.add_paragraph()
final_p.paragraph_format.space_before = Pt(6)
final_p.paragraph_format.space_after = Pt(0)
r1 = final_p.add_run("最终效果：")
set_run_font(r1, size=10.5, bold=True, color=DARK)
r2 = final_p.add_run("网络正常时几乎无感；掉线后持续重试；认证不确定时节流使用浏览器回退；学校计划断网和手动暂停期间不误操作；出现明确凭据或人工验证问题时停止自动提交并等待用户处理。")
set_run_font(r2, size=10.5)

while doc.paragraphs and not doc.paragraphs[-1].text.strip():
    trailing = doc.paragraphs[-1]._element
    trailing.getparent().remove(trailing)

# Save
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUT))
print("DOCX:", OUT)
print("Paragraphs:", len(doc.paragraphs))
print("Tables:", len(doc.tables))
